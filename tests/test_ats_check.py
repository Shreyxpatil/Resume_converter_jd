import pytest
from bs4 import BeautifulSoup
from docx import Document

import src.ats_check as ats_check
import src.docx_converter_general as docx_gen
import src.docx_converter_salesforce as docx_sf
from src.file_extract import extract_text_from_docx_path
from tests.template_fixtures import render_general_html, render_salesforce_html

SAMPLE_HTML = """
<html><body><div class="resume-container" data-template="general">
<table class="resume-table"><tr><td class="left-col">Name</td><td class="right-col"><strong class="name">Jane Doe</strong></td></tr></table>
<table class="projects-table"><tr class="entry-block" data-section-id="work_experience_0">
    <td>Project 1</td>
    <td><div class="project-heading">Project : <span class="entry-title">Acme Corp</span></div>
    <div class="role-heading"><strong>Role:</strong> <span class="entry-role">Engineer</span></div>
    <div class="duration-heading"><strong>Duration:</strong> <span class="entry-duration">Jan 2020 - Dec 2021</span></div>
    <ul><li>Built a service handling 10k requests/sec</li></ul></td>
</tr></table>
</div></body></html>
"""

TEMPLATE_CASES = [
    pytest.param(render_general_html, docx_gen, "jane doe", "acme payments", "built a service", id="general"),
    pytest.param(render_salesforce_html, docx_sf, "raj kumar", "project 1", "built apex triggers", id="salesforce"),
]


@pytest.mark.parametrize("render, converter, name, title, bullet", TEMPLATE_CASES)
def test_real_docx_has_no_parseability_issues(tmp_path, render, converter, name, title, bullet):
    # Both company formats keep every row (identity, skills, projects) in
    # tables in document order — so a linear re-extraction of the real
    # exported DOCX must still find the name first and each project's title
    # before its own bullets.
    html_path, html = render(tmp_path)
    docx_path = str(tmp_path / "out.docx")
    converter.convert_resume_to_docx(html_path, docx_path)

    assert ats_check.check_docx_parseability(docx_path, html) == []


@pytest.mark.parametrize("render, converter, name, title, bullet", TEMPLATE_CASES)
def test_real_docx_preserves_correct_reading_order(tmp_path, render, converter, name, title, bullet):
    html_path, _html = render(tmp_path)
    docx_path = str(tmp_path / "out.docx")
    converter.convert_resume_to_docx(html_path, docx_path)

    text = extract_text_from_docx_path(docx_path).lower()
    name_pos, title_pos, bullet_pos = text.find(name), text.find(title), text.find(bullet)

    assert -1 not in (name_pos, title_pos, bullet_pos)
    assert name_pos < title_pos < bullet_pos


@pytest.mark.parametrize("render, converter, name, title, bullet", TEMPLATE_CASES)
def test_every_compact_level_renders(tmp_path, render, converter, name, title, bullet):
    html_path, _html = render(tmp_path)
    for level in range(len(converter.COMPACT_LEVELS)):
        docx_path = str(tmp_path / f"out_{level}.docx")
        converter.convert_resume_to_docx(html_path, docx_path, level)
        assert len(Document(docx_path).tables) >= 2


@pytest.mark.parametrize("render, converter, name, title, bullet", TEMPLATE_CASES)
def test_bolded_bullet_terms_become_bold_runs(tmp_path, render, converter, name, title, bullet):
    # keyword_highlighter wraps skill mentions in <strong>; the DOCX must
    # keep that as a bold run rather than flattening it to plain text.
    html_path, html = render(tmp_path)
    soup = BeautifulSoup(html, "html.parser")
    li = soup.find("li")
    first_word = li.get_text().split()[0]
    li.string = ""
    strong = soup.new_tag("strong")
    strong.string = first_word
    li.append(strong)
    li.append(" did something measurable")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(str(soup))

    docx_path = str(tmp_path / "out.docx")
    converter.convert_resume_to_docx(html_path, docx_path)
    runs = [r for t in Document(docx_path).tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs]
    assert any(r.bold and r.text.strip() == first_word for r in runs)
    assert any("did something measurable" in r.text and not r.bold for r in runs)


def test_salesforce_converter_raises_instead_of_swallowing_errors(tmp_path):
    with pytest.raises(Exception):
        docx_sf.convert_salesforce_resume(str(tmp_path / "missing.html"), str(tmp_path / "out.docx"))


def test_check_name_position_flags_name_found_late(monkeypatch):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(SAMPLE_HTML, "html.parser")
    # Simulate the real bug: name only appears near the very end of the extracted text.
    extracted = ("x " * 500) + "jane doe"
    issue = ats_check._check_name_position(soup, extracted.lower(), len(extracted))
    assert issue is not None
    assert issue["severity"] == "high"


def test_check_name_position_ok_when_name_near_top():
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(SAMPLE_HTML, "html.parser")
    extracted = "jane doe " + ("x " * 500)
    issue = ats_check._check_name_position(soup, extracted.lower(), len(extracted))
    assert issue is None


def test_find_entries_with_bullets_before_title_detects_reversal():
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(SAMPLE_HTML, "html.parser")
    # Bullet text appears before the entry title in this fake extracted text.
    extracted = "built a service handling 10k requests/sec ... acme corp"
    disordered = ats_check._find_entries_with_bullets_before_title(soup, extracted.lower())
    assert disordered == ["Acme Corp"]


def test_find_entries_with_bullets_before_title_ok_when_in_order():
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(SAMPLE_HTML, "html.parser")
    extracted = "acme corp ... built a service handling 10k requests/sec"
    disordered = ats_check._find_entries_with_bullets_before_title(soup, extracted.lower())
    assert disordered == []


def test_check_docx_parseability_handles_unreadable_docx_gracefully(tmp_path):
    bad_path = str(tmp_path / "not_a_real_docx.docx")
    with open(bad_path, "w") as f:
        f.write("not a real docx file")

    issues = ats_check.check_docx_parseability(bad_path, SAMPLE_HTML)
    assert len(issues) == 1
    assert issues[0]["category"] == "ats_parseability"
    assert issues[0]["severity"] == "medium"


def test_word_count_parity_flagged_when_extracted_text_much_shorter(monkeypatch):
    monkeypatch.setattr(ats_check, "extract_text_from_docx_path", lambda path: "only a few words here")
    issues = ats_check.check_docx_parseability("irrelevant.docx", SAMPLE_HTML)
    parity_issues = [i for i in issues if "re-extracts to only" in i["message"]]
    assert parity_issues
