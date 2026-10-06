from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup
import os
import re

THEME_COLOR = "0d3d62"
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(PROJECT_ROOT, "download.png")

# Typography presets, from the template's normal look to most-compact. They
# exist so a resume that runs past 2 pages can be shrunk to fit WITHOUT ever
# dropping content — app.py's finalize step tries level 0 first and only
# steps up if the rendered page count still exceeds 2. Level 0 is the
# template's original layout; the last level is a legibility floor.
COMPACT_LEVELS = [
    {"margin": 0.5, "font": 12, "heading": 13, "bar": 14, "cell_pad": 8, "logo_w": 3.0, "logo_pad": 12, "cat_before": 12, "line_after": 2},
    {"margin": 0.45, "font": 11, "heading": 12, "bar": 13, "cell_pad": 5, "logo_w": 2.6, "logo_pad": 8, "cat_before": 8, "line_after": 1},
    {"margin": 0.4, "font": 10, "heading": 11, "bar": 12, "cell_pad": 3, "logo_w": 2.2, "logo_pad": 4, "cat_before": 5, "line_after": 0},
]


def html_to_docx(html_file_path, docx_file_path, compact_level=0):
    preset = COMPACT_LEVELS[max(0, min(compact_level, len(COMPACT_LEVELS) - 1))]

    with open(html_file_path, 'r', encoding='utf-8') as file:
        html_content = file.read()

    soup = BeautifulSoup(html_content, 'html.parser')
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(preset["margin"])
        section.bottom_margin = Inches(preset["margin"])
        section.left_margin = Inches(preset["margin"])
        section.right_margin = Inches(preset["margin"])

    try:
        if os.path.exists(LOGO_PATH):
            logo_para = doc.add_paragraph()
            logo_para.paragraph_format.space_before = Pt(preset["logo_pad"])
            logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = logo_para.add_run()
            run.add_picture(LOGO_PATH, width=Inches(preset["logo_w"]))
            logo_para.paragraph_format.space_after = Pt(preset["logo_pad"])
    except Exception: pass

    # === MAIN TABLE ===
    main_table = soup.find('table', {'class': 'resume-table'})
    if main_table:
        doc_table = doc.add_table(rows=0, cols=2)
        doc_table.style = 'Table Grid'
        doc_table.columns[0].width = Inches(1.875)
        doc_table.columns[1].width = Inches(5.625)
        set_table_borders(doc_table)

        for row in main_table.find_all('tr'):
            cells = row.find_all('td')
            if len(cells) == 2:
                doc_row = doc_table.add_row().cells
                set_cell_content(doc_row[0], cells[0].get_text(strip=True), preset, is_header=True)
                set_cell_content(doc_row[1], "", preset, is_header=False, center_align=True)
                process_right_col_content(doc_row[1], cells[1], preset)

    # === PROJECTS SECTION ===
    project_title_div = soup.find('div', {'class': 'project-title-bar'})
    if not project_title_div: project_title_div = soup.find('div', {'class': 'project-title-bar-text-only'})

    projects_table_html = soup.find('table', {'class': 'projects-table'})

    if projects_table_html:
        if project_title_div:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(preset["logo_pad"])
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(project_title_div.get_text(strip=True))
            run.font.name = 'Arial'
            run.bold = True
            run.font.size = Pt(preset["bar"])
            run.font.color.rgb = RGBColor(255, 255, 255)
            set_run_background(run, THEME_COLOR)

        proj_table = doc.add_table(rows=0, cols=2)
        proj_table.style = 'Table Grid'
        proj_table.columns[0].width = Inches(1.875)
        proj_table.columns[1].width = Inches(5.625)
        set_table_borders(proj_table)

        for row in projects_table_html.find_all('tr'):
            cells = row.find_all('td')
            if len(cells) == 2:
                doc_row = proj_table.add_row().cells

                # Project Left Column -> BLUE HEADER STYLE
                set_cell_content(doc_row[0], cells[0].get_text(strip=True), preset, is_header=True, center_align=True, bold=True)

                set_cell_content(doc_row[1], "", preset, is_header=False, center_align=False)
                process_project_content(doc_row[1], cells[1], preset)

    doc.save(docx_file_path)

def set_run_background(run, color_hex):
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    rPr.append(shd)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_cell_content(cell, text, preset, is_header=False, center_align=False, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(preset["cell_pad"])
    paragraph.paragraph_format.space_after = Pt(preset["cell_pad"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center_align else WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    if is_header:
        set_cell_background(cell, THEME_COLOR)
        run = paragraph.add_run(text)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        run.font.size = Pt(preset["font"])
        run.font.name = 'Arial'
    else:
        if text:
            run = paragraph.add_run(text)
            run.font.size = Pt(preset["font"])
            run.font.name = 'Arial'
            run.font.bold = bold

def set_cell_background(cell, color_hex):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_rich_text_runs(paragraph, element, size, bold=False):
    """Renders an element's mixed text/<strong> children as runs, carrying
    <strong> through as bold — the DOCX half of keyword_highlighter's
    skill-mention bolding (get_text() would flatten it to plain text)."""
    children = list(element.children)
    for i, node in enumerate(children):
        is_strong = getattr(node, 'name', None) in ('strong', 'b')
        text = node.get_text() if getattr(node, 'name', None) else str(node)
        text = re.sub(r'\s+', ' ', text)
        if i == 0:
            text = text.lstrip()
        if i == len(children) - 1:
            text = text.rstrip()
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = bold or is_strong
        run.font.size = Pt(size)
        run.font.name = 'Arial'

def process_right_col_content(cell, html_element, preset):
    expertise_categories = html_element.find_all('div', {'class': 'expertise-category'})

    if expertise_categories:
        for cat in expertise_categories:
            title_div = cat.find('div', {'class': 'section-title'})
            if title_div:
                para = cell.add_paragraph()
                para.paragraph_format.space_before = Pt(preset["cat_before"])
                run = para.add_run(title_div.get_text(strip=True))
                run.bold = True
                run.underline = False
                run.font.name = 'Arial'
                run.font.size = Pt(preset["font"])

            skills = cat.find_all('div', {'class': 'skill-item'})
            for skill in skills:
                para = cell.add_paragraph()
                para.paragraph_format.space_after = Pt(preset["line_after"])

                strong_tag = skill.find('strong')
                if strong_tag:
                    label = strong_tag.get_text(strip=True)
                    run = para.add_run(label + " ")
                    run.bold = True
                    run.font.name = 'Arial'
                    run.font.size = Pt(preset["font"])

                    val = skill.get_text().replace(strong_tag.get_text(), '', 1).strip()
                    run2 = para.add_run(val)
                    run2.font.name = 'Arial'
                    run2.font.size = Pt(preset["font"])
                else:
                    run = para.add_run(skill.get_text(strip=True))
                    run.font.name = 'Arial'
                    run.font.size = Pt(preset["font"])
    else:
        text = html_element.get_text(strip=True)
        if text:
            para = cell.paragraphs[0]
            run = para.add_run(text)
            run.font.size = Pt(preset["font"])
            run.font.name = 'Arial'

def _labelled_paragraph(cell, label, value, preset):
    para = cell.add_paragraph()
    para.paragraph_format.space_after = Pt(preset["line_after"])
    run = para.add_run(label)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(preset["font"])
    run2 = para.add_run(value)
    run2.font.name = 'Arial'
    run2.font.size = Pt(preset["font"])

def process_project_content(cell, html_element, preset):
    cell.text = ""

    # 1. Project Heading
    heading = html_element.find('div', {'class': 'project-heading'})
    if heading:
        para = cell.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        raw = heading.get_text(" ", strip=True).replace("Project :", "").replace("Project:", "").strip()
        run = para.add_run(f"Project : {raw}")
        run.bold = True
        run.font.size = Pt(preset["heading"])
        run.font.name = 'Arial'

    # 2. Role
    role = html_element.find('div', {'class': 'role-heading'})
    if role:
        val = role.get_text(" ", strip=True).replace("Role:", "", 1).strip()
        _labelled_paragraph(cell, "Role: ", val, preset)

    # 3. Duration
    dur = html_element.find('div', {'class': 'duration-heading'})
    if dur:
        val = dur.get_text(" ", strip=True).replace("Duration:", "", 1).strip()
        _labelled_paragraph(cell, "Duration: ", val, preset)

    # 4. Links & Description
    link = html_element.find('a')
    if link:
        para = cell.add_paragraph()
        r1 = para.add_run("Link- ")
        r2 = para.add_run(link.get('href', ''))
        for r in (r1, r2):
            r.font.name = 'Arial'
            r.font.size = Pt(preset["font"])

    ul = html_element.find('ul')
    if ul:
        for li in ul.find_all('li'):
            para = cell.add_paragraph()
            para.style = 'List Bullet'
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.space_after = Pt(preset["line_after"])
            add_rich_text_runs(para, li, preset["font"])

    tech = html_element.find('div', {'class': 'techstack'})
    if tech:
        para = cell.add_paragraph()
        para.paragraph_format.space_before = Pt(preset["cell_pad"])
        strong = tech.find('strong')
        if strong:
            r = para.add_run(strong.get_text(strip=True))
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(preset["font"])
            val = tech.get_text().replace(strong.get_text(), '', 1).strip()
            if val:
                r2 = para.add_run("\n" + val)
                r2.font.name = 'Arial'
                r2.font.size = Pt(preset["font"])

def convert_resume_to_docx(html_file_path, docx_file_path=None, compact_level=0):
    if docx_file_path is None:
        docx_file_path = os.path.splitext(html_file_path)[0] + '.docx'
    try:
        html_to_docx(html_file_path, docx_file_path, compact_level)
    except Exception as e:
        print(f"Conversion failed: {e}")
        raise
