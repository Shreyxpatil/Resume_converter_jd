from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup, Tag
import os
import re

# === CONFIGURATION ===
# Word shading fill expects hex without '#'
THEME_COLOR = "1C4587"  # Dark Navy Blue

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(PROJECT_ROOT, "download.png")

# PATH to your badges folder
BADGE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "certificates_badges")

# MAPPING: Keyword in Resume -> Image Filename
CERT_BADGE_MAP = {
    "administrator": "admin_badge.png",
    "platform developer": "pd1_badge.png",
    "app builder": "app_builder_badge.png",
    "consultant": "consultant_badge.png",
    "omnistudio": "omni.png" 
}


def matching_badges(cert_text):
    """Badge image filenames for the certifications named in cert_text, in
    CERT_BADGE_MAP order — shared by the DOCX header and the preview header."""
    cert_text = (cert_text or "").lower()
    return [image for keyword, image in CERT_BADGE_MAP.items() if keyword in cert_text]

# Typography presets, from the template's normal look to most-compact. They
# exist so a resume that runs past 2 pages can be shrunk to fit WITHOUT ever
# dropping content — app.py's finalize step tries level 0 first and only
# steps up if the rendered page count still exceeds 2. Level 0 is the
# template's original layout; the last level is a legibility floor.
COMPACT_LEVELS = [
    {"margin": 0.5, "font": 12, "bar": 14, "cell_pad": 8, "badge_w": 0.85, "logo_w": 2.5, "gap": 6, "line_after": 2},
    {"margin": 0.45, "font": 11, "bar": 13, "cell_pad": 5, "badge_w": 0.75, "logo_w": 2.2, "gap": 4, "line_after": 1},
    {"margin": 0.4, "font": 10, "bar": 12, "cell_pad": 3, "badge_w": 0.65, "logo_w": 2.0, "gap": 2, "line_after": 0},
]

def html_to_docx(html_file_path, docx_file_path, compact_level=0):
    preset = COMPACT_LEVELS[max(0, min(compact_level, len(COMPACT_LEVELS) - 1))]

    with open(html_file_path, 'r', encoding='utf-8') as file:
        html_content = file.read()
    
    soup = BeautifulSoup(html_content, 'html.parser')
    doc = Document()
    
    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(preset["margin"])
        section.bottom_margin = Inches(preset["margin"])
        section.left_margin = Inches(preset["margin"])
        section.right_margin = Inches(preset["margin"])
    
    # === 1. EXTRACT CERTIFICATIONS TEXT ===
    cert_text = ""
    rows = soup.find_all('tr')
    for row in rows:
        cells = row.find_all('td')
        if len(cells) > 0 and "Certifications" in cells[0].get_text():
            if len(cells) > 1:
                cert_text = cells[1].get_text(strip=True).lower()
            break

    # === 2. HEADER: BADGES (Left) + LOGO (Right) ===
    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.0)
    header_table.columns[1].width = Inches(3.5)
    
    # -- A. Add Badges to Left Cell --
    left_cell = header_table.cell(0, 0)
    left_para = left_cell.paragraphs[0]
    left_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    for image_file in matching_badges(cert_text):
        full_badge_path = os.path.join(BADGE_DIR, image_file)
        if os.path.exists(full_badge_path):
            run = left_para.add_run()
            run.add_picture(full_badge_path, width=Inches(preset["badge_w"]))
            run.add_text("  ") # Spacer
        else:
            print(f"⚠️ Warning: Badge image not found at {full_badge_path}")
    
    # -- B. Add Company Logo to Right Cell --
    right_cell = header_table.cell(0, 1)
    right_para = right_cell.paragraphs[0]
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    if os.path.exists(LOGO_PATH):
        run = right_para.add_run()
        run.add_picture(LOGO_PATH, width=Inches(preset["logo_w"]))

    doc.add_paragraph().paragraph_format.space_after = Pt(preset["gap"])

    # === 3. MAIN INFO TABLE ===
    main_table = soup.find('table', {'class': 'resume-table'})
    if main_table:
        t = doc.add_table(rows=0, cols=2)
        t.style = 'Table Grid'
        t.columns[0].width = Inches(1.875)
        t.columns[1].width = Inches(5.625)
        set_table_borders(t)
        
        for row in main_table.find_all('tr'):
            cells = row.find_all('td')
            if len(cells) == 2:
                doc_row = t.add_row().cells
                left_text = cells[0].get_text(strip=True)
                
                # LEFT CELL
                set_cell_content(doc_row[0], left_text, preset, is_header=True, center_align=False)
                
                # RIGHT CELL Logic
                is_name_cell = "name-value-cell" in cells[1].get('class', []) or left_text == "Name"
                
                if is_name_cell:
                    # Name: Blue BG, Centered
                    set_cell_content(doc_row[1], "", preset, is_header=True, center_align=True)
                    process_html_content(doc_row[1], cells[1], preset, text_color="white")
                else:
                    # Standard: White BG, Centered
                    set_cell_content(doc_row[1], "", preset, is_header=False, center_align=True)
                    process_html_content(doc_row[1], cells[1], preset)

    # === 4. SEPARATE PROJECT HEADER ===
    proj_div = soup.find('div', {'class': 'project-title-bar-text-only'})
    if not proj_div: proj_div = soup.find('div', {'class': 'project-title-bar'})
    
    if proj_div:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(preset["cell_pad"] + 4)
        p.paragraph_format.space_after = Pt(4)
        
        run = p.add_run(proj_div.get_text(strip=True))
        run.font.name = 'Arial'
        run.bold = True
        run.font.size = Pt(preset["bar"])
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_run_background(run, THEME_COLOR)

    # === 5. PROJECTS TABLE ===
    proj_table = soup.find('table', {'class': 'projects-table'})
    if proj_table:
        t = doc.add_table(rows=0, cols=2)
        t.style = 'Table Grid'
        t.columns[0].width = Inches(1.875)
        t.columns[1].width = Inches(5.625)
        set_table_borders(t)
        
        for row in proj_table.find_all('tr'):
            cells = row.find_all('td')
            if len(cells) == 2:
                doc_row = t.add_row().cells
                
                # Left Column: Project Name
                set_cell_content(doc_row[0], cells[0].get_text(strip=True), preset, is_header=False, center_align=True, italic=True, bold=True)
                
                # Right Column: Details
                set_cell_content(doc_row[1], "", preset, is_header=False)
                process_project_content(doc_row[1], cells[1], preset)
    
    doc.save(docx_file_path)

# === HELPER FUNCTIONS ===

def set_run_background(run, color_hex):
    """Applies shading (background color) to a specific text run."""
    rPr = run._r.get_or_add_rPr()
    fill = (color_hex or "").replace("#", "")
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    rPr.append(shd)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_cell_content(cell, text, preset, is_header=False, center_align=False, italic=False, bold=False):
    font_size = preset["font"]
    cell.text = ""
    paragraph = cell.paragraphs[0]
    
    # Increase Padding for height (Deepak Singh style)
    paragraph.paragraph_format.space_before = Pt(preset["cell_pad"])
    paragraph.paragraph_format.space_after = Pt(preset["cell_pad"])
    
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center_align else WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    
    # Note: Setting vertical_alignment adds a w:vAlign element to tcPr.
    # Our background color must come BEFORE w:vAlign in the XML for Teams compatibility.
    
    if is_header:
        set_cell_background(cell, THEME_COLOR)
        run = paragraph.add_run(text)
        run.font.name = 'Arial'
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        run.font.size = Pt(font_size)
    else:
        if text:
            run = paragraph.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(font_size)
            run.font.italic = italic
            run.font.bold = bold

def set_cell_background(cell, color_hex):
    """
    Sets cell background color.
    CRITICAL FIX for Teams/Word Online: Ensure <w:shd> is inserted BEFORE <w:vAlign>.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    fill = (color_hex or "").replace("#", "")
    
    # Define shading element
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    
    # Remove any existing shading to prevent duplication
    for existing in tcPr.findall(qn('w:shd')):
        tcPr.remove(existing)
        
    # FIX: Insert <w:shd> before <w:vAlign> or other property tags that must come after it.
    # Common tags that follow 'shd' in OOXML schema: noWrap, tcMar, textDirection, tcFitText, vAlign, hideMark
    
    inserted = False
    for tag in ['w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark']:
        element = tcPr.find(qn(tag))
        if element is not None:
            element.addprevious(shd)
            inserted = True
            break
            
    if not inserted:
        # If none of the later tags exist, safe to append
        tcPr.append(shd)

def process_html_content(cell, html_element, preset, text_color=None):
    text = html_element.get_text(strip=True)
    if text:
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(preset["cell_pad"])
        para.paragraph_format.space_after = Pt(preset["cell_pad"])
        
        run = para.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(preset["font"])
        if text_color == "white":
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True

def add_rich_text_runs(paragraph, element, size, prefix=""):
    """Renders an element's mixed text/<strong> children as runs, carrying
    <strong> through as bold — the DOCX half of keyword_highlighter's
    skill-mention bolding (get_text() would flatten it to plain text)."""
    if prefix:
        run = paragraph.add_run(prefix)
        run.font.name = 'Arial'
        run.font.size = Pt(size)
    children = list(element.children)
    for i, node in enumerate(children):
        is_strong = getattr(node, 'name', None) in ('strong', 'b')
        text = node.get_text() if getattr(node, 'name', None) else str(node)
        text = re.sub(r'\s+', ' ', text)
        if i == 0:
            text = text.lstrip()
        if i == len(children) - 1:
            text = text.rstrip()
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = is_strong
        run.font.name = 'Arial'
        run.font.size = Pt(size)

def process_project_content(cell, html_element, preset):
    font_size = preset["font"]
    cell.text = "" 
    first_element = True
    
    for element in html_element.children:
        if isinstance(element, Tag):
            if element.name == 'div':
                if first_element:
                    para = cell.paragraphs[0]
                    first_element = False
                else:
                    para = cell.add_paragraph()
                
                para.paragraph_format.space_after = Pt(preset["line_after"])
                strong = element.find('strong')
                if strong:
                    label = strong.get_text(strip=True)
                    run = para.add_run(label + " ")
                    run.font.name = 'Arial'
                    run.bold = True
                    
                    val = element.get_text().replace(strong.get_text(), '', 1).strip()
                    if val: 
                        r2 = para.add_run(val)
                        r2.font.name = 'Arial'
                        r2.font.size = Pt(font_size)
                else:
                    text_val = element.get_text(strip=True)
                    run = para.add_run(text_val)
                    run.font.name = 'Arial'
                    run.font.size = Pt(font_size)

                    # Make heading bold + spacing
                    if "Project Description" in text_val:
                        run.bold = True
                        para.paragraph_format.space_before = Pt(preset["cell_pad"])
                        para.paragraph_format.space_after = Pt(preset["line_after"])

            elif element.name == 'ol':
                for i, li in enumerate(element.find_all('li'), 1):
                    if first_element:
                        para = cell.paragraphs[0]
                        first_element = False
                    else:
                        para = cell.add_paragraph()
                        
                    # Handle Indentation manually for best compatibility
                    para.paragraph_format.left_indent = Inches(0.35)
                    para.paragraph_format.first_line_indent = Inches(0)
                    try:
                        para.paragraph_format.tab_stops.clear_all()
                        para.paragraph_format.tab_stops.add_tab_stop(Inches(0.55))
                    except Exception:
                        pass

                    add_rich_text_runs(para, li, font_size, prefix=f"{i}.\t")

            elif element.name == 'ul':
                for li in element.find_all('li'):
                    if first_element:
                        para = cell.paragraphs[0]
                        first_element = False
                    else:
                        para = cell.add_paragraph()

                    para.paragraph_format.left_indent = Inches(0.35)
                    para.paragraph_format.first_line_indent = Inches(0)
                    try:
                        para.paragraph_format.tab_stops.clear_all()
                        para.paragraph_format.tab_stops.add_tab_stop(Inches(0.55))
                    except Exception:
                        pass

                    add_rich_text_runs(para, li, font_size, prefix="•\t")

            elif element.name == 'p':
                if first_element:
                    para = cell.paragraphs[0]
                    first_element = False
                else:
                    para = cell.add_paragraph()
                    
                para.paragraph_format.space_before = Pt(3)
                run = para.add_run(element.get_text(strip=True))
                run.font.name = 'Arial'
                run.font.size = Pt(font_size)

def convert_salesforce_resume(html_file_path, docx_file_path=None, compact_level=0):
    if docx_file_path is None:
        docx_file_path = os.path.splitext(html_file_path)[0] + '.docx'
    try:
        html_to_docx(html_file_path, docx_file_path, compact_level)
    except Exception as e:
        # Used to swallow the error here, so a failed export still reported
        # "completed" with no DOCX on disk — re-raise like the general converter.
        print(f"Conversion failed: {e}")
        raise


# Same name/signature as docx_converter_general so app.py can dispatch either way.
convert_resume_to_docx = convert_salesforce_resume