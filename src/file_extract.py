"""Shared plain-text extraction from uploaded PDF/DOCX files.

Resume extraction and JD extraction used to each carry their own copy of
this logic (one path-based, one bytes-based) — this module covers both
call shapes so there's a single implementation to fix if pypdf/python-docx
extraction behavior ever needs to change.
"""

import io
import os

import docx
from pypdf import PdfReader


def extract_text_from_docx_bytes(file_bytes):
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text_from_pdf_bytes(file_bytes):
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text_from_docx_path(file_path):
    try:
        with open(file_path, "rb") as f:
            return extract_text_from_docx_bytes(f.read())
    except Exception as e:
        raise RuntimeError(f"Failed to read DOCX: {e}")


def extract_text_from_pdf_path(file_path):
    try:
        with open(file_path, "rb") as f:
            return extract_text_from_pdf_bytes(f.read())
    except Exception as e:
        raise RuntimeError(f"Failed to read PDF: {e}")


def extract_text_from_path(file_path):
    """Dispatches on extension for an on-disk resume/JD file (.docx or .pdf)."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return extract_text_from_docx_path(file_path)
    return extract_text_from_pdf_path(file_path)


def extract_text_from_upload(file_bytes, filename):
    """Extracts plain text from an uploaded file's raw bytes (.txt, .docx, .pdf)."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")
    if ext == ".docx":
        return extract_text_from_docx_bytes(file_bytes)
    if ext == ".pdf":
        return extract_text_from_pdf_bytes(file_bytes)
    if ext == ".doc":
        raise ValueError("Legacy .doc format isn't supported — please save as .docx, .pdf, or .txt instead.")

    raise ValueError(f"Unsupported file type: {ext or '(none)'}")
